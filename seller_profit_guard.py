"""CSV-first SKU economics ledger. All monetary computations use Decimal."""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import os
import re
import sqlite3
import sys
import urllib.error
import urllib.parse
import urllib.request
from dataclasses import dataclass
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path

REQUIRED = {"date", "sku", "units", "gross_sales"}
OPTIONAL_TEXT = ("order_id", "order_item_id")
NUMERIC = (
    "units",
    "gross_sales",
    "returns",
    "fees",
    "logistics",
    "ad_spend",
    "ad_sales",
    "cogs",
    "stock",
    "avg_daily_units",
)


def connect(path: str | Path) -> sqlite3.Connection:
    db = sqlite3.connect(path)
    if os.name == "posix" and str(path) != ":memory:":
        try:
            os.chmod(path, 0o600)
        except OSError:
            db.close()
            raise RuntimeError("Could not restrict the database file to the current user") from None
    db.row_factory = sqlite3.Row
    db.executescript("""
    CREATE TABLE IF NOT EXISTS import_batches(id INTEGER PRIMARY KEY, source TEXT NOT NULL, content_hash TEXT NOT NULL UNIQUE,
      imported_at TEXT NOT NULL, accepted INTEGER NOT NULL, rejected INTEGER NOT NULL);
    CREATE TABLE IF NOT EXISTS daily_sku(date TEXT NOT NULL, sku TEXT NOT NULL, units TEXT NOT NULL, gross_sales TEXT NOT NULL,
      returns TEXT NOT NULL, fees TEXT NOT NULL, logistics TEXT NOT NULL, ad_spend TEXT NOT NULL, ad_sales TEXT, cogs TEXT, missing_fields TEXT NOT NULL DEFAULT '',
      stock TEXT, avg_daily_units TEXT, batch_id INTEGER REFERENCES import_batches(id), PRIMARY KEY(date,sku));
    CREATE TABLE IF NOT EXISTS alert_log(id INTEGER PRIMARY KEY, alert_key TEXT NOT NULL, sent_at TEXT NOT NULL, mode TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS products(sku TEXT PRIMARY KEY, first_seen TEXT NOT NULL, last_seen TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS orders(order_id TEXT PRIMARY KEY, order_date TEXT NOT NULL, batch_id INTEGER NOT NULL REFERENCES import_batches(id));
    CREATE TABLE IF NOT EXISTS order_items(
      order_id TEXT NOT NULL REFERENCES orders(order_id), item_key TEXT NOT NULL, sku TEXT NOT NULL REFERENCES products(sku),
      units TEXT NOT NULL, gross_sales TEXT NOT NULL, batch_id INTEGER NOT NULL REFERENCES import_batches(id), source_row INTEGER NOT NULL,
      PRIMARY KEY(order_id,item_key));
    CREATE TABLE IF NOT EXISTS daily_returns(
      date TEXT NOT NULL, sku TEXT NOT NULL REFERENCES products(sku), amount TEXT NOT NULL, batch_id INTEGER NOT NULL REFERENCES import_batches(id),
      PRIMARY KEY(date,sku,batch_id));
    CREATE TABLE IF NOT EXISTS advertising_daily(
      date TEXT NOT NULL, sku TEXT NOT NULL REFERENCES products(sku), spend TEXT, attributed_sales TEXT,
      batch_id INTEGER NOT NULL REFERENCES import_batches(id), PRIMARY KEY(date,sku,batch_id));
    CREATE TABLE IF NOT EXISTS marketplace_fees_daily(
      date TEXT NOT NULL, sku TEXT NOT NULL REFERENCES products(sku), amount TEXT NOT NULL, batch_id INTEGER NOT NULL REFERENCES import_batches(id),
      PRIMARY KEY(date,sku,batch_id));
    CREATE TABLE IF NOT EXISTS logistics_costs_daily(
      date TEXT NOT NULL, sku TEXT NOT NULL REFERENCES products(sku), amount TEXT NOT NULL, batch_id INTEGER NOT NULL REFERENCES import_batches(id),
      PRIMARY KEY(date,sku,batch_id));
    CREATE TABLE IF NOT EXISTS inventory_snapshots(
      date TEXT NOT NULL, sku TEXT NOT NULL REFERENCES products(sku), quantity TEXT NOT NULL, avg_daily_units TEXT,
      batch_id INTEGER NOT NULL REFERENCES import_batches(id), PRIMARY KEY(date,sku,batch_id));
    CREATE INDEX IF NOT EXISTS idx_sku_date ON daily_sku(sku,date);
    """)
    # Additive migration: backfill only facts actually present in legacy daily rows.
    db.execute(
        "INSERT OR IGNORE INTO products(sku,first_seen,last_seen) "
        "SELECT sku,min(date),max(date) FROM daily_sku GROUP BY sku"
    )
    db.execute(
        "INSERT OR IGNORE INTO daily_returns(date,sku,amount,batch_id) "
        "SELECT date,sku,returns,batch_id FROM daily_sku WHERE batch_id IS NOT NULL "
        "AND (','||COALESCE(missing_fields,'')||',') NOT LIKE '%,returns,%'"
    )
    db.execute(
        "INSERT OR IGNORE INTO advertising_daily(date,sku,spend,attributed_sales,batch_id) "
        "SELECT date,sku,CASE WHEN (','||COALESCE(missing_fields,'')||',') LIKE '%,ad_spend,%' THEN NULL ELSE ad_spend END,"
        "ad_sales,batch_id FROM daily_sku WHERE batch_id IS NOT NULL "
        "AND ((','||COALESCE(missing_fields,'')||',') NOT LIKE '%,ad_spend,%' OR ad_sales IS NOT NULL)"
    )
    db.execute(
        "INSERT OR IGNORE INTO marketplace_fees_daily(date,sku,amount,batch_id) "
        "SELECT date,sku,fees,batch_id FROM daily_sku WHERE batch_id IS NOT NULL "
        "AND (','||COALESCE(missing_fields,'')||',') NOT LIKE '%,fees,%'"
    )
    db.execute(
        "INSERT OR IGNORE INTO logistics_costs_daily(date,sku,amount,batch_id) "
        "SELECT date,sku,logistics,batch_id FROM daily_sku WHERE batch_id IS NOT NULL "
        "AND (','||COALESCE(missing_fields,'')||',') NOT LIKE '%,logistics,%'"
    )
    db.execute(
        "INSERT OR IGNORE INTO inventory_snapshots(date,sku,quantity,avg_daily_units,batch_id) "
        "SELECT date,sku,stock,avg_daily_units,batch_id FROM daily_sku "
        "WHERE batch_id IS NOT NULL AND stock IS NOT NULL"
    )
    db.commit()
    return db


def money(value: str | None, field: str, required: bool = False) -> Decimal | None:
    if value is None or value.strip() == "":
        if required:
            raise ValueError(f"{field} is required")
        return None
    try:
        result = Decimal(value.strip().replace(" ", "").replace(",", "."))
    except InvalidOperation as exc:
        raise ValueError(f"{field} must be a number") from exc
    if not result.is_finite():
        raise ValueError(f"{field} must be finite")
    return result


def parse_column_mappings(values: list[str] | None) -> dict[str, str]:
    mappings: dict[str, str] = {}
    allowed = set(NUMERIC) | set(OPTIONAL_TEXT) | {"date", "sku"}
    for value in values or []:
        canonical, separator, source = value.partition("=")
        canonical, source = canonical.strip().lower(), source.strip()
        if not separator or canonical not in allowed or not source:
            raise ValueError("Column mapping must be canonical=source-header (for example gross_sales=Sales)")
        if canonical in mappings:
            raise ValueError(f"Column {canonical} is mapped more than once")
        mappings[canonical] = source
    return mappings


def _validate_rows(
    fieldnames: list[str], records, column_map: dict[str, str] | None = None
) -> tuple[list[dict[str, str]], list[dict[str, object]]]:
    normalized_headers = [name.strip().casefold() for name in fieldnames if name and name.strip()]
    if len(normalized_headers) != len(set(normalized_headers)):
        raise ValueError("Source file has duplicate column headers; rename the columns before import")
    source_headers = {name.strip().casefold(): name for name in fieldnames if name and name.strip()}
    mapping = {
        key: source_headers[key.casefold()]
        for key in (set(NUMERIC) | set(OPTIONAL_TEXT) | {"date", "sku"})
        if key.casefold() in source_headers
    }
    for canonical, source in (column_map or {}).items():
        actual = source_headers.get(source.casefold())
        if actual is None:
            raise ValueError(f"Mapped source header not found: {source}")
        mapping[canonical] = actual
    if len(set(mapping.values())) != len(mapping):
        raise ValueError("Each source header can map to only one canonical column")
    missing = REQUIRED - set(mapping)
    if missing:
        raise ValueError("Missing required columns: " + ", ".join(sorted(missing)))
    normalized = []
    errors = []
    for line, row in enumerate(records, 2):
        item = {key: (row.get(original) or "").strip() for key, original in mapping.items()}
        item["_source_row"] = str(line)
        try:
            d = date.fromisoformat(item["date"])
            if d > date.today():
                raise ValueError("date cannot be in the future")
            if not item["sku"]:
                raise ValueError("sku is required")
            if item.get("order_item_id") and not item.get("order_id"):
                raise ValueError("order_item_id requires order_id")
            for key in NUMERIC:
                value = money(item.get(key), key, key in REQUIRED)
                if value is not None and key in ("units", "returns", "stock", "avg_daily_units") and value < 0:
                    raise ValueError(f"{key} cannot be negative")
                if value is not None and key in ("gross_sales", "fees", "logistics", "ad_spend", "ad_sales", "cogs") and value < 0:
                    raise ValueError(f"{key} cannot be negative")
                if key == "returns" and value is not None and value > money(item["gross_sales"], "gross_sales", True):
                    raise ValueError("returns cannot exceed gross_sales")
            normalized.append(item)
        except ValueError as exc:
            errors.append({"line": line, "error": str(exc), "row": item})
    seen = set()
    duplicates = []
    unique = []
    for item in normalized:
        if item.get("order_id"):
            key = (item["date"], item["sku"], item["order_id"], item.get("order_item_id", ""))
        else:
            key = (item["date"], item["sku"])
        if key in seen:
            duplicates.append({"line_key": key, "error": "duplicate date+sku row"})
        else:
            seen.add(key)
            unique.append(item)
    return unique, errors + duplicates


def parse_csv(
    path: str | Path, column_map: dict[str, str] | None = None
) -> tuple[list[dict[str, str]], list[dict[str, object]]]:
    raw = Path(path).read_bytes()
    if len(raw) > 20 * 1024 * 1024:
        raise ValueError("File exceeds 20 MiB")
    reader = csv.DictReader(io.StringIO(raw.decode("utf-8-sig")))
    if not reader.fieldnames:
        raise ValueError("CSV needs a header row")
    return _validate_rows(reader.fieldnames, reader, column_map)


def parse_xlsx(
    path: str | Path, column_map: dict[str, str] | None = None
) -> tuple[list[dict[str, str]], list[dict[str, object]]]:
    file_path = Path(path)
    if file_path.stat().st_size > 20 * 1024 * 1024:
        raise ValueError("File exceeds 20 MiB")
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("XLSX import requires optional dependency: pip install openpyxl") from exc
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            raise ValueError("XLSX needs a header row")
        names = [str(value).strip() if value is not None else "" for value in header]

        def records():
            for row in rows:
                values = {}
                for index, name in enumerate(names):
                    value = row[index] if index < len(row) else None
                    if isinstance(value, datetime):
                        value = value.date().isoformat()
                    elif isinstance(value, date):
                        value = value.isoformat()
                    values[name] = "" if value is None else str(value)
                yield values

        return _validate_rows(names, records(), column_map)
    finally:
        workbook.close()


def parse_file(
    path: str | Path, column_map: dict[str, str] | None = None
) -> tuple[list[dict[str, str]], list[dict[str, object]]]:
    suffix = Path(path).suffix.lower()
    if suffix == ".csv":
        return parse_csv(path, column_map)
    if suffix == ".xlsx":
        return parse_xlsx(path, column_map)
    raise ValueError("Supported import formats are .csv and .xlsx")


def preview(path: str | Path, column_map: dict[str, str] | None = None) -> dict[str, object]:
    rows, errors = parse_file(path, column_map)
    return {
        "file": str(path),
        "valid_rows": len(rows),
        "rejected_rows": len(errors),
        "sample": [{key: value for key, value in row.items() if key != "_source_row"} for row in rows[:5]],
        "errors": errors[:30],
    }


def _aggregate_import_rows(
    rows: list[dict[str, str]],
) -> tuple[list[tuple[dict[str, str], list[dict[str, str]]]], list[dict[str, object]]]:
    grouped: dict[tuple[str, str], list[dict[str, str]]] = {}
    for row in rows:
        grouped.setdefault((row["date"], row["sku"]), []).append(row)
    output = []
    errors = []
    for (day, sku), items in grouped.items():
        if len(items) == 1:
            output.append((dict(items[0]), items))
            continue
        if not all(item.get("order_id") for item in items):
            errors.extend(
                {"line": item.get("_source_row"), "error": "daily SKU rows cannot be mixed or repeated without explicit order_id"}
                for item in items
            )
            continue
        daily = dict(items[0])
        conflict = False
        for key in NUMERIC:
            present = [item[key] for item in items if item.get(key, "") != ""]
            if key in ("stock", "avg_daily_units"):
                if len(set(present)) > 1:
                    errors.extend(
                        {"line": item.get("_source_row"), "error": f"conflicting daily snapshot values for {key} at {day}/{sku}"}
                        for item in items
                    )
                    conflict = True
                    break
                daily[key] = present[0] if present else ""
            elif len(present) != len(items):
                daily[key] = ""
            else:
                daily[key] = str(sum((money(value, key, True) for value in present), Decimal(0)))
        if conflict:
            continue
        daily["order_id"] = ""
        daily["order_item_id"] = ""
        daily["_source_row"] = items[0].get("_source_row")
        output.append((daily, items))
    return output, errors


def _record_observed_facts(
    db: sqlite3.Connection, batch_id: int, daily: dict[str, str], source_rows: list[dict[str, str]]
) -> None:
    day, sku = daily["date"], daily["sku"]
    db.execute(
        "INSERT INTO products(sku,first_seen,last_seen) VALUES(?,?,?) "
        "ON CONFLICT(sku) DO UPDATE SET first_seen=min(first_seen,excluded.first_seen),last_seen=max(last_seen,excluded.last_seen)",
        (sku, day, day),
    )
    if daily.get("returns", "") != "":
        db.execute("INSERT INTO daily_returns VALUES(?,?,?,?)", (day, sku, daily["returns"], batch_id))
    if daily.get("ad_spend", "") != "" or daily.get("ad_sales", "") != "":
        db.execute(
            "INSERT INTO advertising_daily VALUES(?,?,?,?,?)",
            (day, sku, daily.get("ad_spend") or None, daily.get("ad_sales") or None, batch_id),
        )
    if daily.get("fees", "") != "":
        db.execute("INSERT INTO marketplace_fees_daily VALUES(?,?,?,?)", (day, sku, daily["fees"], batch_id))
    if daily.get("logistics", "") != "":
        db.execute("INSERT INTO logistics_costs_daily VALUES(?,?,?,?)", (day, sku, daily["logistics"], batch_id))
    if daily.get("stock", "") != "":
        db.execute(
            "INSERT INTO inventory_snapshots VALUES(?,?,?,?,?)",
            (day, sku, daily["stock"], daily.get("avg_daily_units") or None, batch_id),
        )
    for source_row in source_rows:
        order_id = source_row.get("order_id", "")
        if not order_id:
            continue
        existing = db.execute("SELECT order_date FROM orders WHERE order_id=?", (order_id,)).fetchone()
        if existing and existing["order_date"] != day:
            raise ValueError(f"Order ID {order_id} has inconsistent dates; correct the source file before import")
        db.execute("INSERT OR IGNORE INTO orders(order_id,order_date,batch_id) VALUES(?,?,?)", (order_id, day, batch_id))
        item_key = source_row.get("order_item_id") or source_row["sku"]
        existing_item = db.execute(
            "SELECT sku,units,gross_sales FROM order_items WHERE order_id=? AND item_key=?", (order_id, item_key)
        ).fetchone()
        values = (source_row["sku"], source_row["units"], source_row["gross_sales"])
        if existing_item and tuple(existing_item) != values:
            raise ValueError(f"Order item {order_id}/{item_key} conflicts with an earlier import")
        db.execute(
            "INSERT OR IGNORE INTO order_items(order_id,item_key,sku,units,gross_sales,batch_id,source_row) VALUES(?,?,?,?,?,?,?)",
            (order_id, item_key, *values, batch_id, int(source_row["_source_row"])),
        )


def _conflicting_order_ids(db: sqlite3.Connection, rows: list[dict[str, str]]) -> set[str]:
    bad: set[str] = set()
    dates: dict[str, str] = {}
    items: dict[tuple[str, str], tuple[str, str, str]] = {}
    for row in rows:
        order_id = row.get("order_id", "")
        if not order_id:
            continue
        if order_id in dates and dates[order_id] != row["date"]:
            bad.add(order_id)
        dates[order_id] = row["date"]
        existing_order = db.execute("SELECT order_date FROM orders WHERE order_id=?", (order_id,)).fetchone()
        if existing_order and existing_order["order_date"] != row["date"]:
            bad.add(order_id)
        item_key = row.get("order_item_id") or row["sku"]
        item_identity = (order_id, item_key)
        item_values = (row["sku"], row["units"], row["gross_sales"])
        if item_identity in items and items[item_identity] != item_values:
            bad.add(order_id)
        items[item_identity] = item_values
        existing_item = db.execute(
            "SELECT sku,units,gross_sales FROM order_items WHERE order_id=? AND item_key=?", item_identity
        ).fetchone()
        if existing_item and tuple(existing_item) != item_values:
            bad.add(order_id)
    return bad


def import_csv(
    db: sqlite3.Connection, path: str | Path, column_map: dict[str, str] | None = None
) -> dict[str, int | bool]:
    raw = Path(path).read_bytes()
    normalized_map = {
        canonical.strip().casefold(): source.strip().casefold() for canonical, source in (column_map or {}).items()
    }
    mapping_bytes = json.dumps(normalized_map, sort_keys=True, separators=(",", ":")).encode()
    digest = hashlib.sha256(raw + b"\0" + mapping_bytes).hexdigest()
    if db.execute("SELECT 1 FROM import_batches WHERE content_hash=?", (digest,)).fetchone():
        return {"duplicate_batch": True, "accepted": 0, "rejected": 0}
    rows, errors = parse_file(path, column_map)
    bad_orders = _conflicting_order_ids(db, rows)
    daily_groups, aggregation_errors = _aggregate_import_rows(rows)
    errors.extend(aggregation_errors)
    if bad_orders:
        kept_groups = []
        for daily, source_rows in daily_groups:
            if any(row.get("order_id") in bad_orders for row in source_rows):
                errors.extend(
                    {"line": row.get("_source_row"), "error": f"Order ID {row['order_id']} conflicts with another row or earlier import"}
                    for row in source_rows
                )
            else:
                kept_groups.append((daily, source_rows))
        daily_groups = kept_groups
    db.execute("SAVEPOINT import_batch")
    try:
        return _persist_import_batch(db, path, digest, errors, daily_groups)
    except BaseException:
        # Unexpected failures must not leave an import batch or a prefix of its rows behind.
        db.execute("ROLLBACK TO SAVEPOINT import_batch")
        db.execute("RELEASE SAVEPOINT import_batch")
        raise


def _persist_import_batch(
    db: sqlite3.Connection,
    path: str | Path,
    digest: str,
    errors: list[dict],
    daily_groups: list[tuple[dict, list[dict]]],
) -> dict[str, int | bool]:
    cur = db.execute(
        "INSERT INTO import_batches(source,content_hash,imported_at,accepted,rejected) VALUES(?,?,?,?,?)",
        (str(path), digest, datetime.now(UTC).isoformat(timespec="seconds"), 0, len(errors)),
    )
    batch = cur.lastrowid
    accepted = 0
    for x, source_rows in daily_groups:
        db.execute("SAVEPOINT import_daily_sku")
        try:
            missing = ",".join(k for k in ("returns", "fees", "logistics", "ad_spend", "cogs") if not x.get(k))
            db.execute(
                "INSERT INTO daily_sku(date,sku,units,gross_sales,returns,fees,logistics,ad_spend,ad_sales,cogs,missing_fields,stock,avg_daily_units,batch_id) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (
                    x["date"],
                    x["sku"],
                    x["units"],
                    x["gross_sales"],
                    x.get("returns") or "0",
                    x.get("fees") or "0",
                    x.get("logistics") or "0",
                    x.get("ad_spend") or "0",
                    x.get("ad_sales") or None,
                    x.get("cogs") or None,
                    missing,
                    x.get("stock") or None,
                    x.get("avg_daily_units") or None,
                    batch,
                ),
            )
            _record_observed_facts(db, batch, x, source_rows)
            db.execute("RELEASE SAVEPOINT import_daily_sku")
            accepted += len(source_rows)
        except sqlite3.IntegrityError:
            db.execute("ROLLBACK TO SAVEPOINT import_daily_sku")
            db.execute("RELEASE SAVEPOINT import_daily_sku")
            errors.extend(
                {"line": source.get("_source_row"), "line_key": (x["date"], x["sku"]), "error": "already imported date+sku"}
                for source in source_rows
            )
        except ValueError as exc:
            db.execute("ROLLBACK TO SAVEPOINT import_daily_sku")
            db.execute("RELEASE SAVEPOINT import_daily_sku")
            errors.extend(
                {"line": source.get("_source_row"), "line_key": (x["date"], x["sku"]), "error": str(exc)}
                for source in source_rows
            )
    db.execute("UPDATE import_batches SET accepted=?,rejected=? WHERE id=?", (accepted, len(errors), batch))
    result = {"duplicate_batch": False, "accepted": accepted, "rejected": len(errors)}
    db.execute("RELEASE SAVEPOINT import_batch")
    return result


@dataclass(frozen=True)
class Metrics:
    gross_sales: Decimal
    returns: Decimal
    net_sales: Decimal | None
    return_rate: Decimal | None
    fees: Decimal
    logistics: Decimal
    ad_spend: Decimal
    ad_sales: Decimal | None
    roas: Decimal | None
    acos: Decimal | None
    cogs: Decimal | None
    contribution_margin: Decimal | None
    margin_complete: bool


def calculate(rows: list[sqlite3.Row]) -> Metrics:
    def total(key: str) -> Decimal:
        return sum((Decimal(str(row[key] or "0")) for row in rows), Decimal(0))

    gross = total("gross_sales")
    returns = total("returns")
    fees = total("fees")
    logistics = total("logistics")
    ads = total("ad_spend")
    missing = set()
    for row in rows:
        missing.update(x for x in (row["missing_fields"] or "").split(",") if x)
    net = None if "returns" in missing else gross - returns
    return_rate = None if "returns" in missing or gross == 0 else returns / gross
    cogs = None if "cogs" in missing else total("cogs")
    ad_sales = None if any(r["ad_sales"] is None for r in rows) else total("ad_sales")
    roas = None if ad_sales is None or ads == 0 else ad_sales / ads
    acos = None if ad_sales in (None, Decimal(0)) else ads / ad_sales
    margin = (
        None
        if (cogs is None or net is None or missing & {"fees", "logistics", "ad_spend"})
        else net - fees - logistics - ads - cogs
    )
    return Metrics(
        gross, returns, net, return_rate, fees, logistics, ads, ad_sales, roas, acos, cogs, margin, margin is not None
    )


def dashboard(db: sqlite3.Connection) -> list[dict[str, object]]:
    out = []
    for skurow in db.execute("SELECT DISTINCT sku FROM daily_sku ORDER BY sku"):
        rows = db.execute("SELECT * FROM daily_sku WHERE sku=? ORDER BY date", (skurow[0],)).fetchall()
        m = calculate(rows)
        days = [date.fromisoformat(r["date"]) for r in rows]
        stockrows = [r for r in rows if r["stock"] is not None and r["avg_daily_units"] not in (None, "0")]
        risk = None
        if stockrows:
            last = stockrows[-1]
            risk = Decimal(last["stock"]) / Decimal(last["avg_daily_units"])
        out.append(
            {
                "sku": skurow[0],
                "first_date": min(days).isoformat(),
                "last_date": max(days).isoformat(),
                "metrics": m,
                "inventory_days": risk,
                "inventory_risk": risk is not None and risk < Decimal(7),
            }
        )
    return out


def evaluate_alerts(
    db: sqlite3.Connection,
    minimum_margin: Decimal,
    min_inventory_days: Decimal,
    cooldown_hours: int = 24,
    dry_run: bool = True,
    minimum_roas: Decimal | None = None,
    maximum_acos: Decimal | None = None,
    trend_days: int = 7,
) -> list[str]:
    if cooldown_hours < 0:
        raise ValueError("Cooldown hours cannot be negative")
    if trend_days < 1:
        raise ValueError("Trend period must be at least one day")
    messages = []
    now = datetime.now(UTC)
    for item in dashboard(db):
        m = item["metrics"]
        sku = item["sku"]
        candidates = []
        if m.contribution_margin is not None and m.contribution_margin < minimum_margin:
            candidates.append(("margin", f"{sku}: contribution margin {m.contribution_margin} below {minimum_margin}"))
        if item["inventory_days"] is not None and item["inventory_days"] < min_inventory_days:
            candidates.append(("stockout", f"{sku}: estimated stock cover {item['inventory_days']:.1f} days"))
        if minimum_roas is not None and m.roas is not None and m.roas < minimum_roas:
            candidates.append(("roas", f"{sku}: ROAS {m.roas:.2f} below {minimum_roas}"))
        if maximum_acos is not None and m.acos is not None and m.acos > maximum_acos:
            candidates.append(("acos", f"{sku}: ACoS {m.acos:.2%} above {maximum_acos:.2%}"))
        all_rows = db.execute("SELECT * FROM daily_sku WHERE sku=? ORDER BY date", (sku,)).fetchall()
        anchor = max(date.fromisoformat(row["date"]) for row in all_rows)
        current_start = anchor - timedelta(days=trend_days - 1)
        previous_start = anchor - timedelta(days=2 * trend_days - 1)
        previous_rows = [row for row in all_rows if previous_start <= date.fromisoformat(row["date"]) < current_start]
        current_rows = [row for row in all_rows if current_start <= date.fromisoformat(row["date"]) <= anchor]
        if previous_rows and current_rows:
            previous_metrics, current_metrics = calculate(previous_rows), calculate(current_rows)
            if (previous_metrics.roas is not None and current_metrics.roas is not None
                    and current_metrics.roas < previous_metrics.roas):
                candidates.append(("roas", (
                    f"{sku}: ROAS declined from {previous_metrics.roas:.2f} "
                    f"to {current_metrics.roas:.2f} across consecutive {trend_days}-day periods"
                )))
            if (previous_metrics.acos is not None and current_metrics.acos is not None
                    and current_metrics.acos > previous_metrics.acos):
                candidates.append(("acos", (
                    f"{sku}: ACoS increased from {previous_metrics.acos:.2%} "
                    f"to {current_metrics.acos:.2%} across consecutive {trend_days}-day periods"
                )))
        seen_candidate_keys: set[str] = set()
        for key, msg in candidates:
            alert = f"{key}:{sku}"
            if alert in seen_candidate_keys:
                continue
            seen_candidate_keys.add(alert)
            last = db.execute(
                "SELECT sent_at FROM alert_log WHERE alert_key=? ORDER BY id DESC LIMIT 1", (alert,)
            ).fetchone()
            if last and (now - datetime.fromisoformat(last[0])).total_seconds() < cooldown_hours * 3600:
                continue
            messages.append(msg)
            if not dry_run:
                db.execute(
                    "INSERT INTO alert_log(alert_key,sent_at,mode) VALUES(?,?,?)",
                    (alert, now.isoformat(timespec="seconds"), "logged"),
                )
    # A batch with rejected rows is actionable even when its valid rows were imported.
    # Key by immutable batch id so cooldowns do not merge separate uploads of the same file.
    for batch in db.execute(
        "SELECT id,source,rejected FROM import_batches WHERE rejected>0 ORDER BY id"
    ):
        alert = f"import_error:batch-{batch['id']}"
        last = db.execute(
            "SELECT sent_at FROM alert_log WHERE alert_key=? ORDER BY id DESC LIMIT 1", (alert,)
        ).fetchone()
        if last and (now - datetime.fromisoformat(last[0])).total_seconds() < cooldown_hours * 3600:
            continue
        messages.append(
            f"Import batch #{batch['id']} from {Path(batch['source']).name} rejected {batch['rejected']} rows"
        )
        if not dry_run:
            db.execute(
                "INSERT INTO alert_log(alert_key,sent_at,mode) VALUES(?,?,?)",
                (alert, now.isoformat(timespec="seconds"), "logged"),
            )
    return messages


def alert_key_for_message(message: str) -> str:
    """Recover a stable cooldown key for either SKU or import-batch alert text."""
    batch = re.match(r"Import batch #(\d+) from ", message)
    if batch:
        return f"import_error:batch-{batch.group(1)}"
    identity, _, detail = message.partition(":")
    if "contribution margin" in detail:
        kind = "margin"
    elif "stock cover" in detail:
        kind = "stockout"
    elif "ROAS" in detail:
        kind = "roas"
    else:
        kind = "acos"
    return f"{kind}:{identity}"


def send_telegram_alerts(
    db: sqlite3.Connection, messages: list[str], token: str, chat_id: str, timeout: float = 10
) -> tuple[int, int]:
    """Explicit, no-retry Telegram delivery. Ambiguous failures are cooldown-logged to avoid duplicate sends."""
    if not token.strip() or not chat_id.strip():
        raise ValueError("Set TELEGRAM_BOT_TOKEN and TELEGRAM_CHAT_ID before --send-telegram")
    sent = 0
    failed = 0
    for message in messages:
        alert_key = alert_key_for_message(message)
        url = f"https://api.telegram.org/bot{token}/sendMessage"
        request = urllib.request.Request(
            url,
            data=urllib.parse.urlencode({"chat_id": chat_id, "text": f"Seller Profit Guard alert\n{message}"}).encode(),
            headers={"Content-Type": "application/x-www-form-urlencoded"},
        )
        outcome = "sent"
        try:
            with urllib.request.urlopen(request, timeout=timeout) as response:
                payload = json.loads(response.read(100_000))
            if not isinstance(payload, dict) or not payload.get("ok"):
                outcome = "unknown"
        except (urllib.error.URLError, TimeoutError, json.JSONDecodeError, OSError):
            # Do not retry this non-idempotent request automatically. A timeout may occur after delivery.
            outcome = "unknown"
        if outcome == "sent":
            sent += 1
        else:
            failed += 1
        db.execute(
            "INSERT INTO alert_log(alert_key,sent_at,mode) VALUES(?,?,?)",
            (alert_key, datetime.now(UTC).isoformat(timespec="seconds"), outcome),
        )
    return sent, failed


def serializable(item: dict[str, object]) -> dict[str, object]:
    result = {k: v for k, v in item.items() if k != "metrics"}
    result["metrics"] = {k: (str(v) if isinstance(v, Decimal) else v) for k, v in vars(item["metrics"]).items()}
    result["inventory_days"] = str(item["inventory_days"]) if item["inventory_days"] is not None else None
    return result


def spreadsheet_safe_text(value: str) -> str:
    """Neutralize formula-leading SKU values in CSV exports opened by spreadsheet apps."""
    if value.lstrip(" \t\r\n").startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def export_csv(db: sqlite3.Connection, path: str | Path) -> None:
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["date", "sku", *NUMERIC])
        for row in db.execute("SELECT * FROM daily_sku ORDER BY date,sku"):
            writer.writerow([row["date"], spreadsheet_safe_text(row["sku"]), *[row[k] for k in NUMERIC]])
    if os.name == "posix":
        os.chmod(path, 0o600)


def fact_summary(db: sqlite3.Connection) -> dict[str, int]:
    tables = (
        "products", "orders", "order_items", "daily_returns", "advertising_daily",
        "marketplace_fees_daily", "logistics_costs_daily", "inventory_snapshots", "import_batches",
    )
    return {table: int(db.execute(f"SELECT count(*) FROM {table}").fetchone()[0]) for table in tables}


def main() -> int:
    p = argparse.ArgumentParser(description="Marketplace unit economics, CSV-first")
    p.add_argument("--db", default="seller_profit.sqlite3")
    sub = p.add_subparsers(dest="cmd", required=True)
    x = sub.add_parser("preview")
    x.add_argument("file")
    x.add_argument("--map", dest="mappings", action="append", default=[], metavar="CANONICAL=HEADER")
    x = sub.add_parser("import")
    x.add_argument("file")
    x.add_argument("--map", dest="mappings", action="append", default=[], metavar="CANONICAL=HEADER")
    sub.add_parser("dashboard")
    sub.add_parser("imports")
    sub.add_parser("facts", help="Show counts for imported normalized facts and order rows")
    x = sub.add_parser("alerts")
    x.add_argument("--min-margin", default="0")
    x.add_argument("--min-stock-days", default="7")
    x.add_argument("--min-roas")
    x.add_argument("--max-acos")
    x.add_argument("--cooldown-hours", type=int, default=24)
    x.add_argument("--trend-days", type=int, default=7)
    group = x.add_mutually_exclusive_group()
    group.add_argument("--commit-log", action="store_true")
    group.add_argument("--send-telegram", action="store_true")
    x = sub.add_parser("export")
    x.add_argument("file")
    args = p.parse_args()
    try:
        if args.cmd == "preview":
            print(json.dumps(preview(args.file, parse_column_mappings(args.mappings)), ensure_ascii=False, indent=2))
            return 0
        db = connect(args.db)
        if args.cmd == "import":
            print(json.dumps(import_csv(db, args.file, parse_column_mappings(args.mappings))))
            db.commit()
        elif args.cmd == "imports":
            print(json.dumps([dict(row) for row in db.execute(
                "SELECT id,source,imported_at,accepted,rejected FROM import_batches ORDER BY id DESC"
            )], ensure_ascii=False, indent=2))
        elif args.cmd == "facts":
            print(json.dumps(fact_summary(db), ensure_ascii=False, indent=2))
        elif args.cmd == "dashboard":
            print(json.dumps([serializable(x) for x in dashboard(db)], indent=2))
        elif args.cmd == "alerts":
            alerts = evaluate_alerts(
                db,
                Decimal(args.min_margin),
                Decimal(args.min_stock_days),
                cooldown_hours=args.cooldown_hours,
                dry_run=True,
                minimum_roas=Decimal(args.min_roas) if args.min_roas is not None else None,
                maximum_acos=Decimal(args.max_acos) if args.max_acos is not None else None,
                trend_days=args.trend_days,
            )
            if args.send_telegram:
                sent, failed = send_telegram_alerts(
                    db, alerts, os.environ.get("TELEGRAM_BOT_TOKEN", ""), os.environ.get("TELEGRAM_CHAT_ID", "")
                )
                print(
                    f"Telegram delivery: {sent} sent, {failed} unknown/failed; requests are never retried automatically"
                )
            else:
                for a in alerts:
                    print(("DRY-RUN: " if not args.commit_log else "LOGGED: ") + a)
            if not alerts:
                print("No alerts")
            if args.commit_log:
                # This mode records a local preview only; it never sends external messages.
                for a in alerts:
                    db.execute(
                        "INSERT INTO alert_log(alert_key,sent_at,mode) VALUES(?,?,?)",
                        (alert_key_for_message(a), datetime.now(UTC).isoformat(timespec="seconds"), "logged"),
                    )
            db.commit()
        elif args.cmd == "export":
            export_csv(db, args.file)
        db.close()
        return 0
    except (ValueError, OSError, UnicodeError, InvalidOperation) as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
